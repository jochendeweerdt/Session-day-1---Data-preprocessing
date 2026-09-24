"""Point the notebooks and the slide deck to the real GitHub repository.

Usage (from the repository root):
    python tools/set_repo_links.py your-account/your-repo "path/to/Session 1.2 - ... .pptx"

- Notebooks: replaces the placeholder OWNER/REPO (Colab badge and data URL).
- Deck: updates the "open in Colab" hyperlinks, the repository name in the text, and regenerates
  the QR codes (pictures named "QR NB1" ... "QR NB4"). Nothing else in the deck is touched.
Requires: pip install python-pptx qrcode pillow
"""
import glob, io, re, sys
import qrcode
from pptx import Presentation

PLACEHOLDER = "OWNER/REPO"
repo = sys.argv[1]
deck = sys.argv[2] if len(sys.argv) > 2 else None

for nb in glob.glob("notebooks/*.ipynb"):
    text = open(nb, encoding="utf-8").read()
    if PLACEHOLDER in text:
        open(nb, "w", encoding="utf-8").write(text.replace(PLACEHOLDER, repo))
        print("updated", nb)

if deck:
    prs = Presentation(deck)
    n_links = n_qr = 0
    for slide in prs.slides:
        for rel in slide.part.rels.values():
            if rel.is_external and "colab.research.google.com/github/" in rel.target_ref:
                new_url = re.sub(r"(colab\.research\.google\.com/github/)[^/]+/[^/]+/", rf"\g<1>{repo}/", rel.target_ref)
                rel._target = new_url
                rel.__dict__.pop("target_ref", None)      # python-pptx caches this property
                n_links += 1
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    for r in p.runs:
                        if PLACEHOLDER in r.text:
                            r.text = r.text.replace(PLACEHOLDER, repo)
            if shape.shape_type == 13 and shape.name.startswith("QR NB"):
                nb_no = shape.name[-1]
                target = [rel.target_ref for rel in slide.part.rels.values()
                          if rel.is_external and f"/NB{nb_no}_" in rel.target_ref][0]
                buf = io.BytesIO(); qrcode.make(target, border=1, box_size=10).save(buf, format="PNG")
                image_part = slide.part.related_part(shape._element.blipFill.blip.rEmbed)
                image_part._blob = buf.getvalue()
                n_qr += 1
    prs.save(deck)
    print(f"deck: {n_links} links and {n_qr} QR codes updated")
